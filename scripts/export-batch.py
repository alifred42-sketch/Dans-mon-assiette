#!/usr/bin/env python3
"""Ajoute BATCH et COURSES officiels du tableur dans carnet.json."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ORIG = Path("/tmp/orig/sheet.xlsx")
OUT = Path(__file__).resolve().parents[1] / "data" / "carnet.json"
BOX = re.compile(r"^[\s☐☑✅✔●·]+")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return BOX.sub("", str(value).strip())


def main() -> None:
    wb = load_workbook(ORIG, data_only=True)

    batch: list[dict] = []
    ws = wb["_BATCH_DATA"]
    for r in range(2, (ws.max_row or 0) + 1):
        raw_week = ws.cell(r, 1).value
        text = cell_text(ws.cell(r, 2).value)
        kind = cell_text(ws.cell(r, 3).value)
        if raw_week is None or not text:
            continue
        try:
            week = int(raw_week)
        except (TypeError, ValueError):
            continue
        if week < 1 or week > 52 or text.startswith("📅"):
            continue
        batch.append({"week": week, "text": text, "type": kind})

    courses: list[dict] = []
    ws = wb["_COURSES_DATA"]
    for r in range(2, (ws.max_row or 0) + 1):
        raw_week = ws.cell(r, 1).value
        dish = cell_text(ws.cell(r, 4).value)
        ingredient = cell_text(ws.cell(r, 5).value)
        aisle = cell_text(ws.cell(r, 6).value)
        if raw_week is None or not ingredient:
            continue
        try:
            week = int(raw_week)
        except (TypeError, ValueError):
            continue
        if week < 1 or week > 52:
            continue
        courses.append(
            {
                "week": week,
                "dish": dish,
                "ingredient": ingredient,
                "aisle": aisle or "🛒 À vérifier",
            }
        )

    wb.close()
    data = json.loads(OUT.read_text(encoding="utf-8"))
    data["batch"] = batch
    data["courses"] = courses
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        "batch",
        len(batch),
        "week1",
        sum(1 for i in batch if i["week"] == 1),
        "courses",
        len(courses),
        "week1",
        sum(1 for i in courses if i["week"] == 1),
    )


if __name__ == "__main__":
    main()
